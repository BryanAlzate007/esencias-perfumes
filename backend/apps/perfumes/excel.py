import re
import unicodedata
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO

from django.db import transaction
from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Main_chords, Perfume

HEADERS = (
    ("id", "id"),
    ("nombre", "name"),
    ("marca", "brand"),
    ("descripción", "description"),
    ("notas", "notes"),
    ("catálogo", "catalog"),
    ("precio", "price"),
    ("precio_usd", "price_usd"),
    ("gramos", "grams"),
    ("precio_gramo", "gram_price"),
    ("color", "color"),
    ("acordes", "main_chords"),
    ("activo", "is_active"),
)

HEADER_FIELDS = {
    "id": "id",
    "nombre": "name",
    "name": "name",
    "marca": "brand",
    "brand": "brand",
    "descripcion": "description",
    "description": "description",
    "notas": "notes",
    "notes": "notes",
    "catalogo": "catalog",
    "catalog": "catalog",
    "precio": "price",
    "price": "price",
    "precio usd": "price_usd",
    "price usd": "price_usd",
    "gramos": "grams",
    "grams": "grams",
    "precio gramo": "gram_price",
    "precio del gramo": "gram_price",
    "valor del gramo": "gram_price",
    "gram price": "gram_price",
    "color": "color",
    "acordes": "main_chords",
    "acordes principales": "main_chords",
    "main chords": "main_chords",
    "activo": "is_active",
    "active": "is_active",
}

CATALOGS = {
    "caballero": Perfume.Catalog.CABALLERO,
    "hombre": Perfume.Catalog.CABALLERO,
    "men": Perfume.Catalog.CABALLERO,
    "dama": Perfume.Catalog.DAMA,
    "mujer": Perfume.Catalog.DAMA,
    "women": Perfume.Catalog.DAMA,
    "arabe": Perfume.Catalog.ARABE,
    "arabian": Perfume.Catalog.ARABE,
}

MONEY = Decimal("0.01")


class ExcelImportError(Exception):
    def __init__(self, detail, errors=None):
        super().__init__(detail)
        self.detail = detail
        self.errors = errors or []


class RowError(Exception):
    pass


def export_perfumes_workbook(queryset):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Catalogo"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1A140C")

    for column, (label, _field) in enumerate(HEADERS, start=1):
        cell = sheet.cell(1, column, label)
        cell.font = header_font
        cell.fill = header_fill

    for row_index, perfume in enumerate(queryset.prefetch_related("main_chords"), start=2):
        values = {
            "id": perfume.id,
            "name": perfume.name,
            "brand": perfume.brand,
            "description": perfume.description,
            "notes": perfume.notes,
            "catalog": perfume.catalog,
            "price": float(perfume.price),
            "price_usd": float(perfume.price_usd) if perfume.price_usd is not None else None,
            "grams": perfume.grams,
            "gram_price": float(perfume.gram_price),
            "color": perfume.color,
            "main_chords": ", ".join(perfume.main_chords.order_by("name").values_list("name", flat=True)),
            "is_active": "si" if perfume.is_active else "no",
        }
        for column, (_label, field) in enumerate(HEADERS, start=1):
            sheet.cell(row_index, column, values[field])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(sheet.max_row, 1)}"
    for column, (label, _field) in enumerate(HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = max(len(label) + 4, 14)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="catalogo-perfumes.xlsx"'
    workbook.save(response)
    return response


def import_perfumes_workbook(uploaded):
    rows = _read_rows(uploaded)
    if not rows:
        raise ExcelImportError("El archivo no tiene perfumes para cargar.")

    created = 0
    updated = 0
    errors = []
    with transaction.atomic():
        for item in rows:
            try:
                was_created = _upsert(item["fields"])
            except RowError as exc:
                errors.append({"row": item["row"], "message": str(exc)})
                continue
            if was_created:
                created += 1
            else:
                updated += 1
        if errors:
            transaction.set_rollback(True)

    if errors:
        raise ExcelImportError("No se importó el archivo. Revisa las filas indicadas.", errors)
    return {"created": created, "updated": updated}


def _read_rows(uploaded):
    uploaded.seek(0)
    try:
        workbook = load_workbook(BytesIO(uploaded.read()), data_only=True, read_only=True)
    except Exception as exc:
        raise ExcelImportError("El archivo de Excel no es válido.") from exc

    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    try:
        header = next(iterator)
    except StopIteration as exc:
        raise ExcelImportError("El archivo está vacío.") from exc

    columns = {}
    for index, label in enumerate(header or ()):
        if label is None or str(label).strip() == "":
            continue
        field = HEADER_FIELDS.get(_normalize(label))
        if field is None:
            continue
        if field in columns:
            raise ExcelImportError(f"La columna {label} está repetida.")
        columns[field] = index

    if "name" not in columns or "brand" not in columns:
        raise ExcelImportError("El archivo debe incluir las columnas nombre y marca.")

    rows = []
    for offset, raw in enumerate(iterator, start=2):
        if raw is None or all(_is_blank(value) for value in raw):
            continue
        fields = {}
        for field, index in columns.items():
            value = raw[index] if index < len(raw) else None
            if _is_blank(value):
                continue
            fields[field] = value
        if fields:
            rows.append({"row": offset, "fields": fields})
    return rows


def _upsert(fields):
    perfume = _find_perfume(fields)
    if perfume is None:
        payload, chords = _create_payload(fields)
        perfume = Perfume.objects.create(**payload)
        if chords:
            perfume.main_chords.set(chords)
        return True

    changes, chords = _update_payload(fields)
    for key, value in changes.items():
        setattr(perfume, key, value)
    if changes:
        perfume.save()
    if chords is not None:
        perfume.main_chords.set(chords)
    return False


def _find_perfume(fields):
    if "id" in fields:
        perfume_id = _parse_id(fields["id"])
        perfume = Perfume.objects.filter(pk=perfume_id).first()
        if perfume is None:
            raise RowError(f"No existe un perfume con id {perfume_id}.")
        return perfume

    name = _text(fields.get("name"), "El nombre", 200)
    brand = _text(fields.get("brand"), "La marca", 120)
    matches = list(Perfume.objects.filter(name__iexact=name, brand__iexact=brand))
    if len(matches) > 1:
        raise RowError("Hay varios perfumes con el mismo nombre y marca. Incluye el id para actualizarlos.")
    return matches[0] if matches else None


def _create_payload(fields):
    missing = [label for key, label in (
        ("name", "nombre"),
        ("brand", "marca"),
        ("description", "descripción"),
        ("catalog", "catálogo"),
        ("price", "precio"),
    ) if key not in fields]
    if missing:
        raise RowError(f"Faltan datos obligatorios: {', '.join(missing)}.")

    payload = {
        "name": _text(fields["name"], "El nombre", 200),
        "brand": _text(fields["brand"], "La marca", 120),
        "description": _text(fields["description"], "La descripción"),
        "notes": _text(fields["notes"], "Las notas", 200) if "notes" in fields else "",
        "catalog": _parse_catalog(fields["catalog"]),
        "price": _parse_money(fields["price"], "El precio"),
        "price_usd": _parse_money(fields["price_usd"], "El precio USD") if "price_usd" in fields else None,
        "color": _text(fields["color"], "El color", 200) if "color" in fields else "",
        "grams": _parse_grams(fields["grams"]) if "grams" in fields else 100,
        "gram_price": _parse_money(fields["gram_price"], "El precio del gramo") if "gram_price" in fields else Decimal("0.00"),
        "is_active": _parse_bool(fields["is_active"]) if "is_active" in fields else True,
    }
    chords = _parse_chords(fields["main_chords"]) if "main_chords" in fields else []
    return payload, chords


def _update_payload(fields):
    changes = {}
    if "name" in fields:
        changes["name"] = _text(fields["name"], "El nombre", 200)
    if "brand" in fields:
        changes["brand"] = _text(fields["brand"], "La marca", 120)
    if "description" in fields:
        changes["description"] = _text(fields["description"], "La descripción")
    if "notes" in fields:
        changes["notes"] = _text(fields["notes"], "Las notas", 200)
    if "catalog" in fields:
        changes["catalog"] = _parse_catalog(fields["catalog"])
    if "price" in fields:
        changes["price"] = _parse_money(fields["price"], "El precio")
    if "price_usd" in fields:
        changes["price_usd"] = _parse_money(fields["price_usd"], "El precio USD")
    if "color" in fields:
        changes["color"] = _text(fields["color"], "El color", 200)
    if "grams" in fields:
        changes["grams"] = _parse_grams(fields["grams"])
    if "gram_price" in fields:
        changes["gram_price"] = _parse_money(fields["gram_price"], "El precio del gramo")
    if "is_active" in fields:
        changes["is_active"] = _parse_bool(fields["is_active"])
    chords = _parse_chords(fields["main_chords"]) if "main_chords" in fields else None
    return changes, chords


def _parse_id(value):
    if isinstance(value, bool):
        raise RowError("El id no es válido.")
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = str(value).strip()
    if text.isdigit():
        return int(text)
    raise RowError("El id no es válido.")


def _parse_catalog(value):
    catalog = CATALOGS.get(_normalize(value))
    if catalog is None:
        raise RowError("El catálogo debe ser caballero, dama o arabe.")
    return catalog


def _parse_money(value, label):
    try:
        if isinstance(value, Decimal):
            number = value
        elif isinstance(value, bool):
            raise InvalidOperation
        elif isinstance(value, (int, float)):
            number = Decimal(str(value))
        else:
            text = str(value).strip().replace(" ", "").replace("$", "")
            if "," in text and "." in text:
                if text.rfind(",") > text.rfind("."):
                    text = text.replace(".", "").replace(",", ".")
                else:
                    text = text.replace(",", "")
            elif "," in text:
                text = text.replace(",", ".")
            number = Decimal(text)
        number = number.quantize(MONEY, rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError) as exc:
        raise RowError(f"{label} no es un número válido.") from exc
    if number < 0 or number >= Decimal("100000000"):
        raise RowError(f"{label} está fuera de rango.")
    return number


def _parse_grams(value):
    if isinstance(value, bool):
        raise RowError("Los gramos deben ser un número entero.")
    if isinstance(value, int):
        number = value
    elif isinstance(value, float) and value.is_integer():
        number = int(value)
    else:
        text = str(value).strip()
        if not text.isdigit():
            raise RowError("Los gramos deben ser un número entero.")
        number = int(text)
    if number < 1:
        raise RowError("Los gramos deben ser mayores a cero.")
    return number


def _parse_bool(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and int(value) in (0, 1) and float(value) in (0.0, 1.0):
        return bool(int(value))
    key = _normalize(value)
    if key in {"si", "true", "yes", "1", "activo"}:
        return True
    if key in {"no", "false", "0", "inactivo"}:
        return False
    raise RowError("Activo debe ser si o no.")


def _parse_chords(value):
    text = str(value).strip()
    names = [part.strip() for part in re.split(r"[,;]", text) if part.strip()]
    chords = []
    for name in names:
        if len(name) > 200:
            raise RowError("Un acorde supera los 200 caracteres.")
        chord = Main_chords.objects.filter(name__iexact=name).first()
        if chord is None:
            chord = Main_chords.objects.create(name=name, description="")
        chords.append(chord)
    return chords


def _text(value, label, max_length=None):
    if value is None:
        raise RowError(f"{label} es obligatorio.")
    text = str(value).strip()
    if not text:
        raise RowError(f"{label} es obligatorio.")
    if max_length and len(text) > max_length:
        raise RowError(f"{label} supera los {max_length} caracteres.")
    return text


def _normalize(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(character for character in text if not unicodedata.combining(character))
    text = text.replace("_", " ").strip().lower()
    return " ".join(text.split())


def _is_blank(value):
    return value is None or (isinstance(value, str) and not value.strip())
