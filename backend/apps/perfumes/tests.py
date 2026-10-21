import json
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from PIL import Image
from rest_framework.test import APIClient

from apps.accounts.models import Profile
from apps.accounts.models import Profile_person
from apps.perfumes.models import Container, Main_chords, Perfume, PerfumeImage

User = get_user_model()


class PerfumeCatalogTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.admin = User.objects.create_user(
            username="admin",
            email="admin@example.com",
            password="Admin1234!",
            is_staff=True,
        )
        self.admin.profile.role = Profile.Role.ADMIN
        self.admin.profile.save()
        self.chord = Main_chords.objects.create(name="Amaderado", description="Notas de madera")

    def test_admin_can_create_perfume_with_relations(self):
        self.client.force_login(self.admin)
        response = self.client.post(
            "/api/v1/perfumes/",
            data={
                "name": "Aurora",
                "brand": "Esencias",
                "description": "Cítrica",
                "notes": "Bergamota",
                "catalog": "dama",
                "price": "70.00",
                "price_usd": "18.50",
                "color": "Ámbar",
                "main_chords": [self.chord.id],
                "is_active": True,
            },
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.content)
        perfume = Perfume.objects.get(name="Aurora")
        self.assertEqual(perfume.color, "Ámbar")
        self.assertEqual(list(perfume.main_chords.values_list("id", flat=True)), [self.chord.id])
        self.assertEqual(response.json()["main_chords"], [self.chord.id])
        self.assertEqual(response.json()["catalog"], "dama")

    def test_admin_can_upload_several_perfume_images(self):
        self.client.force_login(self.admin)
        created = self.client.post(
            "/api/v1/perfumes/",
            data={
                "name": "Atlas",
                "brand": "Esencias",
                "description": "Árabe",
                "catalog": "arabe",
                "price": "90.00",
                "main_chords": json.dumps([self.chord.id]),
                "images": [jpeg("one.jpg"), jpeg("two.jpg")],
            },
            format="multipart",
        )
        self.assertEqual(created.status_code, 201, created.content)
        perfume = Perfume.objects.get(name="Atlas")
        self.assertEqual(perfume.images.count(), 2)
        self.assertEqual(perfume.images.filter(is_primary=True).count(), 1)
        self.assertEqual(perfume.catalog, Perfume.Catalog.ARABE)

        removed = perfume.images.order_by("order").first()
        updated = self.client.patch(
            f"/api/v1/perfumes/{perfume.id}/",
            data={"remove_image_ids": [removed.id], "images": [jpeg("three.jpg")]},
            format="multipart",
        )
        self.assertEqual(updated.status_code, 200, updated.content)
        self.assertEqual(PerfumeImage.objects.filter(perfume=perfume).count(), 2)
        self.assertFalse(PerfumeImage.objects.filter(id=removed.id).exists())

    def test_admin_can_list_main_chords(self):
        self.client.force_login(self.admin)
        response = self.client.get("/api/v1/main-chords/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()[0]["name"], "Amaderado")

    def test_admin_can_manage_lookup_catalogs(self):
        self.client.force_login(self.admin)
        chord = self.client.post(
            "/api/v1/main-chords/",
            data={"name": "Cítrico", "description": "Notas frescas"},
            format="json",
        )
        self.assertEqual(chord.status_code, 201, chord.content)

        container = self.client.post(
            "/api/v1/containers/",
            data={
                "name": "Frasco 50 ml",
                "description": "Vidrio ámbar",
                "type": "Frasco",
                "size": "50 ml",
                "weight": "120 g",
                "volume": "50 ml",
                "material": "Vidrio",
                "color": "Ámbar",
                "main_chords": [self.chord.id],
                "is_active": True,
            },
            format="json",
        )
        self.assertEqual(container.status_code, 201, container.content)
        self.assertEqual(Container.objects.get(name="Frasco 50 ml").main_chords.count(), 1)

        person = self.client.post(
            "/api/v1/profile-persons/",
            data={"name": "Nocturna", "description": "Prefiere notas intensas"},
            format="json",
        )
        self.assertEqual(person.status_code, 201, person.content)
        self.assertEqual(Profile_person.objects.filter(name="Nocturna").count(), 1)

    def test_admin_can_export_and_import_catalog_excel(self):
        perfume = Perfume.objects.create(
            name="Aurora",
            brand="Esencias",
            description="Cítrica",
            notes="Bergamota",
            catalog=Perfume.Catalog.DAMA,
            price="70.00",
            price_usd="18.50",
            grams=50,
            gram_price="1.20",
            color="Ámbar",
            is_active=True,
        )
        perfume.main_chords.add(self.chord)
        self.client.force_login(self.admin)

        exported = self.client.get("/api/v1/perfumes/export/")
        self.assertEqual(exported.status_code, 200)
        sheet = load_workbook(BytesIO(exported.content)).active
        self.assertEqual(sheet.cell(1, 2).value, "nombre")
        self.assertEqual(sheet.cell(2, 2).value, "Aurora")
        self.assertEqual(sheet.cell(2, 12).value, "Amaderado")

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["nombre", "marca", "descripción", "catálogo", "precio", "acordes", "activo"])
        sheet.append(["Aurora", "Esencias", "Más cítrica", "dama", 80, "Amaderado, Floral", "si"])
        sheet.append(["Noche", "Esencias", "Intensa", "caballero", "90,50", "Especiado", "no"])
        uploaded = self.client.post(
            "/api/v1/perfumes/import/",
            data={"file": workbook_file(workbook)},
            format="multipart",
        )
        self.assertEqual(uploaded.status_code, 200, uploaded.content)
        self.assertEqual(uploaded.json(), {"created": 1, "updated": 1})
        perfume.refresh_from_db()
        self.assertEqual(perfume.price, 80)
        self.assertEqual(perfume.description, "Más cítrica")
        self.assertEqual(set(perfume.main_chords.values_list("name", flat=True)), {"Amaderado", "Floral"})
        created = Perfume.objects.get(name="Noche")
        self.assertEqual(created.catalog, Perfume.Catalog.CABALLERO)
        self.assertEqual(created.price, Decimal("90.50"))
        self.assertFalse(created.is_active)

    def test_invalid_excel_does_not_change_the_catalog(self):
        self.client.force_login(self.admin)
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["nombre", "marca", "descripción", "catálogo", "precio"])
        sheet.append(["Valido", "Esencias", "Texto", "dama", 10])
        sheet.append(["Invalido", "Esencias", "Texto", "otro", 10])
        uploaded = self.client.post(
            "/api/v1/perfumes/import/",
            data={"file": workbook_file(workbook)},
            format="multipart",
        )
        self.assertEqual(uploaded.status_code, 400)
        self.assertEqual(uploaded.json()["errors"][0]["row"], 3)
        self.assertFalse(Perfume.objects.filter(name="Valido").exists())

    def test_customer_cannot_import_or_export_catalog(self):
        customer = User.objects.create_user(username="cliente", email="cliente@example.com", password="Cliente1234!")
        self.client.force_login(customer)
        self.assertEqual(self.client.get("/api/v1/perfumes/export/").status_code, 403)
        self.assertEqual(self.client.post("/api/v1/perfumes/import/", data={}, format="multipart").status_code, 403)

    def test_public_list_pages_twelve_until_filtered(self):
        for index in range(13):
            Perfume.objects.create(
                name=f"Fragancia {index:02d}",
                brand="Esencias",
                description="Notas",
                price="50.00",
                catalog=Perfume.Catalog.DAMA if index < 3 else Perfume.Catalog.CABALLERO,
            )

        listed = self.client.get("/api/v1/perfumes/")
        self.assertEqual(listed.status_code, 200)
        body = listed.json()
        self.assertEqual(body["count"], 13)
        self.assertEqual(len(body["results"]), 12)

        second = self.client.get("/api/v1/perfumes/", {"page": 2})
        self.assertEqual(len(second.json()["results"]), 1)

        filtered = self.client.get("/api/v1/perfumes/", {"catalog": "dama"})
        self.assertIsInstance(filtered.json(), list)
        self.assertEqual(len(filtered.json()), 3)

        searched = self.client.get("/api/v1/perfumes/", {"search": "Fragancia 0"})
        self.assertIsInstance(searched.json(), list)
        self.assertEqual(len(searched.json()), 10)


def workbook_file(workbook):
    buffer = BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(
        "catalogo.xlsx",
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def jpeg(name):
    buffer = BytesIO()
    Image.new("RGB", (4, 4), "red").save(buffer, format="JPEG")
    return SimpleUploadedFile(name, buffer.getvalue(), content_type="image/jpeg")
